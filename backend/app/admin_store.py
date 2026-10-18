from __future__ import annotations

import json
import csv
import logging
import os
import sqlite3
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy import create_engine, text

from .venue_schema import ensure_venue_schema, sync_normalized_venue


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = ROOT / "venue_pipeline_backend" / "data_auto_only_final_v2"
DEFAULT_DB = ROOT / "backend" / "data" / "shuffl_admin.db"
LIVE_GOOGLE_FIELDS = (
    "google_live_cache", "google_live_retrieved_at", "google_rating", "google_reviews",
    "google_address", "google_hours", "latitude", "longitude",
)
_log = logging.getLogger("shuffl.admin_store")


def _load_local_env() -> None:
    """Load local settings for CLI syncs as well as the FastAPI process."""
    for env_path in (ROOT / "backend" / ".env", ROOT / "venue_pipeline_backend" / ".env"):
        if not env_path.exists():
            continue
        for raw_line in env_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip("\"").strip("'")
            if key and key not in os.environ:
                os.environ[key] = value


_load_local_env()


def _database_url() -> str:
    value = os.getenv("DATABASE_URL", "").strip()
    if value:
        return value
    return f"sqlite:///{DEFAULT_DB.as_posix()}"


class AdminStore:
    def __init__(self) -> None:
        self.engine = create_engine(_database_url(), future=True, pool_pre_ping=True)
        self.schema_error: str | None = None
        try:
            self.ensure_schema()
        except Exception as exc:  # keep /health/readiness available when DB is down
            self.schema_error = str(exc)[:240]
            _log.warning("database_schema_unavailable error=%s", self.schema_error)

    def ensure_schema(self) -> None:
        # Create the normalized app domain tables alongside the catalogue
        # tables. Runtime aggregate persistence remains backwards-compatible,
        # while new repositories can move entity-by-entity to these tables.
        from .models import Base
        Base.metadata.create_all(self.engine)
        with self.engine.begin() as conn:
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS admin_venues (
                    venue_id VARCHAR(180) PRIMARY KEY,
                    google_place_id VARCHAR(180) UNIQUE,
                    display_name VARCHAR(255) NOT NULL DEFAULT '',
                    city VARCHAR(120) NOT NULL DEFAULT '',
                    state VARCHAR(120) NOT NULL DEFAULT 'Delhi',
                    area VARCHAR(160) NOT NULL DEFAULT '',
                    admission_status VARCHAR(40) NOT NULL DEFAULT 'Insufficient evidence',
                    verification_status VARCHAR(40) NOT NULL DEFAULT 'in_process',
                    official_url TEXT,
                    google_lookup_url TEXT,
                    data_json TEXT NOT NULL,
                    source_synced_at VARCHAR(40),
                    updated_at VARCHAR(40) NOT NULL
                )
            """))
            try:
                conn.execute(text("ALTER TABLE admin_venues ADD COLUMN state VARCHAR(120) NOT NULL DEFAULT 'Delhi'"))
            except Exception:
                pass
            for statement in (
                "ALTER TABLE admin_venues ADD COLUMN eligibility_status VARCHAR(40) NOT NULL DEFAULT 'NEEDS_REVIEW'",
                "ALTER TABLE admin_venues ADD COLUMN publication_status VARCHAR(40) NOT NULL DEFAULT 'DRAFT'",
            ):
                try:
                    conn.execute(text(statement))
                except Exception:
                    pass
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS admin_users (
                    user_id VARCHAR(180) PRIMARY KEY,
                    email VARCHAR(320) UNIQUE,
                    display_name VARCHAR(180) NOT NULL DEFAULT '',
                    role VARCHAR(40) NOT NULL DEFAULT 'user',
                    status VARCHAR(40) NOT NULL DEFAULT 'active',
                    last_seen_at VARCHAR(40),
                    created_at VARCHAR(40) NOT NULL
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS admin_audit_log (
                    id VARCHAR(64) PRIMARY KEY,
                    venue_id VARCHAR(180) NOT NULL,
                    action VARCHAR(80) NOT NULL,
                    actor VARCHAR(180) NOT NULL DEFAULT 'admin',
                    before_json TEXT,
                    after_json TEXT,
                    created_at VARCHAR(40) NOT NULL
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS admin_settings (
                    setting_key VARCHAR(120) PRIMARY KEY,
                    setting_value TEXT NOT NULL,
                    updated_by VARCHAR(180) NOT NULL DEFAULT 'system',
                    updated_at VARCHAR(40) NOT NULL
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS venue_publication_versions (
                    version_id VARCHAR(64) PRIMARY KEY,
                    venue_id VARCHAR(180) NOT NULL,
                    version_number INTEGER NOT NULL,
                    status VARCHAR(40) NOT NULL,
                    snapshot_json TEXT NOT NULL,
                    created_by VARCHAR(180) NOT NULL,
                    created_at VARCHAR(40) NOT NULL,
                    UNIQUE (venue_id, version_number)
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS venue_manual_overrides (
                    id VARCHAR(64) PRIMARY KEY,
                    venue_id VARCHAR(180) NOT NULL,
                    field_name VARCHAR(120) NOT NULL,
                    automatic_value TEXT,
                    manual_value TEXT,
                    effective_value TEXT,
                    reason TEXT,
                    edited_by VARCHAR(180) NOT NULL,
                    edited_at VARCHAR(40) NOT NULL,
                    reset_at VARCHAR(40)
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS venue_change_history (
                    id VARCHAR(64) PRIMARY KEY,
                    venue_id VARCHAR(180) NOT NULL,
                    action VARCHAR(80) NOT NULL,
                    before_json TEXT,
                    after_json TEXT,
                    actor VARCHAR(180) NOT NULL,
                    created_at VARCHAR(40) NOT NULL
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS venue_pipeline_runs (
                    run_id VARCHAR(64) PRIMARY KEY,
                    status VARCHAR(40) NOT NULL,
                    source_dir TEXT,
                    source_rows INTEGER NOT NULL DEFAULT 0,
                    inserted_rows INTEGER NOT NULL DEFAULT 0,
                    updated_rows INTEGER NOT NULL DEFAULT 0,
                    error_count INTEGER NOT NULL DEFAULT 0,
                    started_at VARCHAR(40) NOT NULL,
                    finished_at VARCHAR(40),
                    summary_json TEXT
                )
            """))
            conn.execute(text("""
                CREATE TABLE IF NOT EXISTS venue_media (
                    media_id VARCHAR(64) PRIMARY KEY,
                    venue_id VARCHAR(180) NOT NULL,
                    url TEXT NOT NULL,
                    media_type VARCHAR(40) NOT NULL DEFAULT 'image',
                    source_type VARCHAR(40) NOT NULL DEFAULT 'admin_upload',
                    rights_status VARCHAR(80) NOT NULL DEFAULT 'unknown',
                    alt_text VARCHAR(255) NOT NULL DEFAULT '',
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    is_cover INTEGER NOT NULL DEFAULT 0,
                    created_by VARCHAR(180) NOT NULL,
                    created_at VARCHAR(40) NOT NULL
                )
            """))
            for statement in (
                "ALTER TABLE venue_media ADD COLUMN crop_json TEXT",
                "ALTER TABLE venue_media ADD COLUMN variants_json TEXT",
            ):
                try:
                    conn.execute(text(statement))
                except Exception:
                    pass
            # Keep the normalized venue platform available for both the
            # dashboard and the public mobile API.  This is additive: the
            # existing admin_venues/data_json records remain the lifecycle
            # root and are never deleted by bootstrap.
            ensure_venue_schema(conn)
            conn.execute(text("UPDATE admin_venues SET eligibility_status=CASE WHEN admission_status='Admitted' THEN 'AUTO_APPROVED' WHEN admission_status='Rejected' THEN 'REJECTED' ELSE 'NEEDS_REVIEW' END WHERE eligibility_status IS NULL OR eligibility_status='' OR eligibility_status='NEEDS_REVIEW'"))
            conn.execute(text("UPDATE admin_venues SET publication_status=CASE WHEN admission_status='Admitted' AND verification_status <> 'rejected' THEN 'PUBLISHED' ELSE 'DRAFT' END WHERE publication_status IS NULL OR publication_status='' OR publication_status='DRAFT'"))

    @staticmethod
    def _json(value: Any) -> str:
        return json.dumps(value, ensure_ascii=False, default=str)

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).isoformat()

    @staticmethod
    def _state_for_city(city: str) -> str:
        city = (city or "").strip().lower()
        if city in {"gurugram", "faridabad", "rewari", "nuh", "palwal", "sonipat", "panipat", "rohtak", "jind", "karnal", "bhiwani", "charkhi dadri", "jhajjar", "mahendragarh"}:
            return "Haryana"
        if city in {"noida", "greater noida", "ghaziabad", "hapur", "meerut", "bulandshahr", "baghpat", "baraut", "muzaffarnagar", "shamli"}:
            return "Uttar Pradesh"
        if city in {"alwar", "bharatpur", "bhiwadi"}:
            return "Rajasthan"
        return "Delhi"

    def _row(self, row: Any) -> dict[str, Any]:
        data = json.loads(row.data_json or "{}")
        # Legacy snapshots may contain Google display payloads. They are not
        # part of the durable API contract and are removed from every response.
        for key in LIVE_GOOGLE_FIELDS:
            data.pop(key, None)
        data.update({
            "venue_id": row.venue_id,
            "google_place_id": row.google_place_id,
            "display_name": row.display_name,
            "city": row.city,
            "state": getattr(row, "state", "Delhi") or "Delhi",
            "area": row.area,
            "admission_status": row.admission_status,
            "verification_status": row.verification_status,
            "eligibility_status": getattr(row, "eligibility_status", "NEEDS_REVIEW") or "NEEDS_REVIEW",
            "publication_status": getattr(row, "publication_status", "DRAFT") or "DRAFT",
            "official_url": row.official_url or "",
            "google_lookup_url": row.google_lookup_url or "",
            "updated_at": row.updated_at,
        })
        return data

    def purge_live_google_cache(self) -> dict[str, int]:
        """Remove legacy Google display payloads from durable lifecycle JSON."""
        keys = LIVE_GOOGLE_FIELDS
        processed = removed = 0
        with self.engine.begin() as conn:
            rows = conn.execute(text("SELECT venue_id, data_json FROM admin_venues")).all()
            for venue_id, raw in rows:
                payload = json.loads(raw or "{}")
                changed = False
                for key in keys:
                    if key in payload:
                        payload.pop(key, None)
                        changed = True
                if changed:
                    conn.execute(text("UPDATE admin_venues SET data_json=:data_json, updated_at=:updated_at WHERE venue_id=:venue_id"), {"data_json": self._json(payload), "updated_at": self._now(), "venue_id": venue_id})
                    removed += 1
                processed += 1
        return {"processed": processed, "removed": removed}

    def sync_from_source(self, source_dir: str | Path | None = None) -> dict[str, int]:
        if os.getenv("VENUE_PIPELINE_LOCKED", "0").strip().lower() in {"1", "true", "yes", "on"}:
            raise RuntimeError("Venue pipeline is locked; unlock VENUE_PIPELINE_LOCKED before syncing")
        source = Path(source_dir or os.getenv("VENUE_SOURCE_DIR", DEFAULT_SOURCE))
        def load(name: str) -> list[dict[str, Any]]:
            path = source / f"{name}.json"
            if not path.exists():
                return []
            payload = json.loads(path.read_text(encoding="utf-8"))
            return payload.get(name, [])

        candidates = load("candidates")
        by_id: dict[str, dict[str, Any]] = {row.get("venue_id", ""): dict(row) for row in candidates if row.get("venue_id")}
        for table in ("venue_identity", "venue_hours", "venue_time_slots", "venue_scores", "venue_tags", "venue_sources", "gold_list"):
            grouped: dict[str, list[dict[str, Any]]] = {}
            for row in load(table):
                grouped.setdefault(row.get("venue_id", ""), []).append(row)
            for venue_id, rows in grouped.items():
                if venue_id in by_id:
                    by_id[venue_id][table] = rows
        identity = {row.get("venue_id"): row for row in load("venue_identity")}
        now = self._now()
        inserted = updated = 0
        with self.engine.begin() as conn:
            for venue_id, data in by_id.items():
                old = conn.execute(text("SELECT data_json, publication_status FROM admin_venues WHERE venue_id=:id"), {"id": venue_id}).fetchone()
                existing = json.loads(old[0]) if old and old[0] else {}
                for key in LIVE_GOOGLE_FIELDS:
                    existing.pop(key, None)
                override = existing.get("admin_override", {})
                merged = {**data, "admin_override": override}
                for key in LIVE_GOOGLE_FIELDS:
                    merged.pop(key, None)
                info = identity.get(venue_id, {})
                cached = existing.get("google_live_cache") or {}
                cached_name = ((cached.get("displayName") or {}).get("text") or "").strip() if isinstance(cached, dict) else ""
                display_name = info.get("official_name") or existing.get("display_name") or cached_name or ""
                values = {
                    "venue_id": venue_id,
                    "google_place_id": data.get("google_place_id", ""),
                    "display_name": display_name,
                    "city": data.get("discovered_city", ""),
                    "state": data.get("discovered_state") or self._state_for_city(data.get("discovered_city", "")),
                    "area": data.get("discovered_area", ""),
                    "admission_status": data.get("admission_decision", "Insufficient evidence"),
                    "verification_status": override.get("verification_status", data.get("verification_status", "in_process")),
                    "eligibility_status": data.get("eligibility_status") or ({"Admitted": "AUTO_APPROVED", "Rejected": "REJECTED"}.get(data.get("admission_decision"), "NEEDS_REVIEW")),
                    "publication_status": old[1] if old and old[1] else ("PUBLISHED" if data.get("admission_decision") == "Admitted" else "DRAFT"),
                    "official_url": info.get("official_url", ""),
                    "google_lookup_url": data.get("google_lookup_url", ""),
                    "data_json": self._json(merged),
                    "source_synced_at": now,
                    "updated_at": now,
                }
                if self.engine.dialect.name == "mysql":
                    conn.execute(text("""
                        INSERT INTO admin_venues
                        (venue_id,google_place_id,display_name,city,state,area,admission_status,verification_status,eligibility_status,publication_status,official_url,google_lookup_url,data_json,source_synced_at,updated_at)
                        VALUES (:venue_id,:google_place_id,:display_name,:city,:state,:area,:admission_status,:verification_status,:eligibility_status,:publication_status,:official_url,:google_lookup_url,:data_json,:source_synced_at,:updated_at)
                        ON DUPLICATE KEY UPDATE
                          google_place_id=VALUES(google_place_id), display_name=VALUES(display_name), city=VALUES(city), state=VALUES(state), area=VALUES(area), admission_status=VALUES(admission_status), verification_status=VALUES(verification_status), eligibility_status=VALUES(eligibility_status), official_url=VALUES(official_url), google_lookup_url=VALUES(google_lookup_url), data_json=VALUES(data_json), source_synced_at=VALUES(source_synced_at), updated_at=VALUES(updated_at)
                    """), values)
                else:
                    conn.execute(text("""
                        INSERT INTO admin_venues
                        (venue_id,google_place_id,display_name,city,state,area,admission_status,verification_status,eligibility_status,publication_status,official_url,google_lookup_url,data_json,source_synced_at,updated_at)
                        VALUES (:venue_id,:google_place_id,:display_name,:city,:state,:area,:admission_status,:verification_status,:eligibility_status,:publication_status,:official_url,:google_lookup_url,:data_json,:source_synced_at,:updated_at)
                        ON CONFLICT(venue_id) DO UPDATE SET
                          google_place_id=excluded.google_place_id, display_name=excluded.display_name,
                          city=excluded.city, state=excluded.state, area=excluded.area, admission_status=excluded.admission_status,
                          verification_status=excluded.verification_status, eligibility_status=excluded.eligibility_status, publication_status=excluded.publication_status, official_url=excluded.official_url,
                          google_lookup_url=excluded.google_lookup_url, data_json=excluded.data_json,
                          source_synced_at=excluded.source_synced_at, updated_at=excluded.updated_at
                    """), values)
                sync_normalized_venue(conn, merged, now)
                if old:
                    updated += 1
                else:
                    inserted += 1
        return {"source_rows": len(by_id), "inserted": inserted, "updated": updated}

    def backfill_normalized(self) -> dict[str, int]:
        """Populate canonical tables from existing lifecycle records."""
        now = self._now()
        processed = failed = 0
        with self.engine.begin() as conn:
            rows = conn.execute(text("""
                SELECT venue_id, display_name, city, state, area, official_url,
                       google_place_id, google_lookup_url, data_json
                FROM admin_venues
            """)).all()
            for venue_id, root_name, root_city, root_state, root_area, root_official_url, root_place_id, root_lookup_url, raw in rows:
                try:
                    payload = json.loads(raw or "{}")
                    payload["venue_id"] = payload.get("venue_id") or venue_id
                    # Older imports kept lifecycle identity on admin_venues
                    # rather than inside data_json. Preserve that canonical
                    # identity during repair without using live Google data.
                    payload.setdefault("display_name", root_name or "")
                    payload.setdefault("discovered_city", root_city or "")
                    payload.setdefault("discovered_state", root_state or "")
                    payload.setdefault("discovered_area", root_area or "")
                    payload.setdefault("official_url", root_official_url or "")
                    payload.setdefault("google_place_id", root_place_id or "")
                    payload.setdefault("google_lookup_url", root_lookup_url or "")
                    if not isinstance(payload.get("admin_override"), dict):
                        payload["admin_override"] = {}
                    sync_normalized_venue(conn, payload, now)
                    processed += 1
                except Exception:
                    failed += 1
        return {"processed": processed, "failed": failed}

    def normalized_venue(self, venue_id: str) -> dict[str, Any]:
        """Read structured venue facts for admin and public consumers."""
        with self.engine.connect() as conn:
            identity = conn.execute(text("SELECT * FROM venue_identity WHERE venue_id=:id"), {"id": venue_id}).mappings().first()
            scores = conn.execute(text("SELECT * FROM venue_scores WHERE venue_id=:id ORDER BY factor"), {"id": venue_id}).mappings().all()
            tags = conn.execute(text("SELECT * FROM venue_tags WHERE venue_id=:id ORDER BY tag"), {"id": venue_id}).mappings().all()
            sources = conn.execute(text("SELECT * FROM venue_sources WHERE venue_id=:id ORDER BY field_name"), {"id": venue_id}).mappings().all()
            slots = conn.execute(text("SELECT * FROM venue_time_slots WHERE venue_id=:id AND is_current=1 ORDER BY day_of_week,start_time"), {"id": venue_id}).mappings().all()
            metros = conn.execute(text("SELECT * FROM venue_metros WHERE venue_id=:id ORDER BY distance_km"), {"id": venue_id}).mappings().all()
        return {
            "identity": dict(identity) if identity else {},
            "scores": [dict(row) for row in scores],
            "tags": [dict(row) for row in tags],
            "sources": [dict(row) for row in sources],
            "time_slots": [dict(row) for row in slots],
            "metros": [dict(row) for row in metros],
        }

    def summary(self) -> dict[str, Any]:
        with self.engine.connect() as conn:
            total = conn.execute(text("SELECT COUNT(*) FROM admin_venues")).scalar() or 0
            by_status = conn.execute(text("SELECT admission_status, COUNT(*) FROM admin_venues GROUP BY admission_status")).all()
            by_verification = conn.execute(text("SELECT verification_status, COUNT(*) FROM admin_venues GROUP BY verification_status")).all()
            by_eligibility = conn.execute(text("SELECT eligibility_status, COUNT(*) FROM admin_venues GROUP BY eligibility_status")).all()
            by_publication = conn.execute(text("SELECT publication_status, COUNT(*) FROM admin_venues GROUP BY publication_status")).all()
            audit_count = conn.execute(text("SELECT COUNT(*) FROM admin_audit_log")).scalar() or 0
        configured_lock = os.getenv("VENUE_PIPELINE_LOCKED", "0").strip().lower() in {"1", "true", "yes", "on"}
        persisted_lock = self.get_setting("pipeline_locked", "1" if configured_lock else "0") == "1"
        return {"venues": total, "users": len(self.users()), "by_admission": dict(by_status), "by_verification": dict(by_verification), "by_eligibility": dict(by_eligibility), "by_publication": dict(by_publication), "audit_events": audit_count, "pipeline_locked": persisted_lock, "database_url": _database_url().split("@")[-1]}

    def list_venues(self, search: str = "", status: str = "", verification: str = "", page: int = 1, page_size: int = 50, city: str = "", state: str = "", area: str = "", venue_type: str = "", publication: str = "") -> dict[str, Any]:
        page = max(1, page); page_size = min(200, max(1, page_size)); offset = (page - 1) * page_size
        clauses = ["1=1"]; params: dict[str, Any] = {"limit": page_size, "offset": offset}
        if search:
            clauses.append("(display_name LIKE :search OR area LIKE :search OR city LIKE :search OR google_place_id LIKE :search)"); params["search"] = f"%{search}%"
        if status:
            clauses.append("admission_status=:status"); params["status"] = status
        if verification:
            clauses.append("verification_status=:verification"); params["verification"] = verification
        if city:
            clauses.append("city=:city"); params["city"] = city
        if state:
            clauses.append("state=:state"); params["state"] = state
        if area:
            clauses.append("area LIKE :area"); params["area"] = f"%{area}%"
        if venue_type:
            clauses.append("data_json LIKE :venue_type"); params["venue_type"] = f'%"discovered_type": "{venue_type}"%'
        if publication:
            if publication not in {"DRAFT", "PUBLISHED", "UNPUBLISHED", "ARCHIVED"}:
                raise ValueError("Invalid publication status")
            clauses.append("publication_status=:publication"); params["publication"] = publication
        where = " AND ".join(clauses)
        with self.engine.connect() as conn:
            total = conn.execute(text(f"SELECT COUNT(*) FROM admin_venues WHERE {where}"), params).scalar() or 0
            rows = conn.execute(text(f"SELECT * FROM admin_venues WHERE {where} ORDER BY updated_at DESC LIMIT :limit OFFSET :offset"), params).fetchall()
        return {"items": [self._row(row) for row in rows], "page": page, "page_size": page_size, "total": total}

    def get_venue(self, venue_id: str) -> dict[str, Any] | None:
        with self.engine.connect() as conn:
            row = conn.execute(text("SELECT * FROM admin_venues WHERE venue_id=:id"), {"id": venue_id}).fetchone()
        return self._row(row) if row else None

    def create_venue(self, payload: dict[str, Any], actor: str = "admin") -> dict[str, Any]:
        now = self._now()
        venue_id = str(payload.get("venue_id") or f"manual:{uuid4().hex}")
        data = dict(payload)
        data["venue_id"] = venue_id
        data.setdefault("admission_decision", "Insufficient evidence")
        data.setdefault("verification_status", "in_process")
        values = {
            "venue_id": venue_id,
            "google_place_id": payload.get("google_place_id") or None,
            "display_name": payload.get("display_name") or payload.get("name") or "",
            "city": payload.get("city") or "New Delhi",
            "state": payload.get("state") or "Delhi",
            "area": payload.get("area") or "",
            "admission_status": payload.get("admission_status") or "Insufficient evidence",
            "verification_status": payload.get("verification_status") or "in_process",
            "eligibility_status": payload.get("eligibility_status") or "NEEDS_REVIEW",
            "publication_status": "DRAFT",
            "official_url": payload.get("official_url") or payload.get("website_url") or "",
            "google_lookup_url": payload.get("google_lookup_url") or "",
            "data_json": self._json(data),
            "source_synced_at": now,
            "updated_at": now,
        }
        with self.engine.begin() as conn:
            conn.execute(text("""
                INSERT INTO admin_venues
                (venue_id,google_place_id,display_name,city,state,area,admission_status,verification_status,eligibility_status,publication_status,official_url,google_lookup_url,data_json,source_synced_at,updated_at)
                VALUES (:venue_id,:google_place_id,:display_name,:city,:state,:area,:admission_status,:verification_status,:eligibility_status,:publication_status,:official_url,:google_lookup_url,:data_json,:source_synced_at,:updated_at)
            """), values)
            conn.execute(text("INSERT INTO admin_audit_log (id,venue_id,action,actor,after_json,created_at) VALUES (:id,:venue_id,:action,:actor,:after_json,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "action": "manual_create", "actor": actor, "after_json": self._json(data), "created_at": now})
            sync_normalized_venue(conn, data, now)
        return self.get_venue(venue_id) or data

    def import_rows(self, rows: list[dict[str, Any]], actor: str = "import") -> dict[str, int]:
        created = updated = skipped = 0
        for raw in rows:
            row = {str(key).strip(): value for key, value in raw.items() if key is not None and str(key).strip()}
            if not any(row.get(key) for key in ("display_name", "name", "official_name", "venue_name")):
                skipped += 1; continue
            place_id = row.get("google_place_id") or row.get("place_id")
            existing = None
            if place_id:
                with self.engine.connect() as conn:
                    existing_row = conn.execute(text("SELECT venue_id FROM admin_venues WHERE google_place_id=:place_id"), {"place_id": str(place_id)}).fetchone()
                    existing = existing_row[0] if existing_row else None
            payload = dict(row)
            payload.update({"google_place_id": place_id, "display_name": row.get("display_name") or row.get("name") or row.get("official_name") or row.get("venue_name"), "venue_id": existing or row.get("venue_id")})
            if existing:
                current = self.get_venue(existing) or {}
                self.update_venue(existing, payload, actor)
                updated += 1
            else:
                self.create_venue(payload, actor); created += 1
        return {"created": created, "updated": updated, "skipped": skipped, "rows": len(rows)}

    def distinct_values(self) -> dict[str, list[str]]:
        with self.engine.connect() as conn:
            def values(column: str) -> list[str]:
                return [str(row[0]) for row in conn.execute(text(f"SELECT DISTINCT {column} FROM admin_venues WHERE {column} IS NOT NULL AND {column} <> '' ORDER BY {column}" )).all()]
            return {"cities": values("city"), "states": values("state"), "areas": values("area")}

    def area_catalog(self) -> list[dict[str, Any]]:
        """Return the configured area-list plus live venue counts.

        The area list is the discovery source of truth; database-only areas are
        appended so manually imported records remain discoverable in admin.
        """
        configured: dict[tuple[str, str], dict[str, Any]] = {}
        area_file = ROOT / "venue_pipeline_backend" / "config" / "areas.csv"
        if area_file.exists():
            with area_file.open(newline="", encoding="utf-8-sig") as handle:
                for row in csv.DictReader(handle):
                    area = (row.get("area") or "").strip()
                    city = (row.get("city") or "").strip()
                    if area:
                        configured[(area.casefold(), city.casefold())] = {"area": area, "city": city, "state": (row.get("state") or "").strip(), "source": "area_list"}
        with self.engine.connect() as conn:
            rows = conn.execute(text("SELECT area, city, state, COUNT(*) AS venue_count FROM admin_venues WHERE area IS NOT NULL AND area <> '' GROUP BY area, city, state ORDER BY area, city")).mappings().all()
        for row in rows:
            key = (str(row["area"]).casefold(), str(row["city"] or "").casefold())
            item = configured.setdefault(key, {"area": str(row["area"]), "city": str(row["city"] or ""), "state": str(row["state"] or ""), "source": "database"})
            item["venue_count"] = int(row["venue_count"] or 0)
        return sorted(configured.values(), key=lambda item: (str(item.get("city", "")), str(item.get("area", ""))))

    def update_venue(self, venue_id: str, patch: dict[str, Any], actor: str = "admin") -> dict[str, Any]:
        allowed = {"display_name", "city", "state", "area", "official_url", "google_lookup_url", "google_place_id", "discovered_type", "website_url", "menu_url", "booking_url", "social_urls", "metro_name", "metro_distance_km", "photos", "verification_status", "eligibility_status", "operator_class", "fit_out_status", "fit_out_year", "manual_notes", "manual_vibe", "manual_tags", "manual_opening_hours", "manual_time_slots", "published", "override_reason"}
        clean = {key: value for key, value in patch.items() if key in allowed}
        if "verification_status" in clean and clean["verification_status"] not in {"verified", "rejected", "in_process", "update_required"}:
            raise ValueError("Invalid verification_status")
        if "eligibility_status" in clean and clean["eligibility_status"] not in {"AUTO_APPROVED", "PROVISIONAL", "NEEDS_REVIEW", "REJECTED"}:
            raise ValueError("Invalid eligibility_status")
        current = self.get_venue(venue_id)
        if not current:
            return None
        before = dict(current)
        override = dict(current.get("admin_override") or {})
        override.update(clean)
        now = self._now()
        updated_data = {**current, **{key: value for key, value in clean.items() if key not in {"verification_status"}}, "admin_override": override}
        sql_values = {"id": venue_id, "display_name": updated_data.get("display_name", current.get("display_name", "")), "city": updated_data.get("city", current.get("city", "")), "state": updated_data.get("state", current.get("state", "Delhi")), "area": updated_data.get("area", current.get("area", "")), "google_place_id": updated_data.get("google_place_id") or None, "official_url": updated_data.get("official_url", current.get("official_url", "")), "google_lookup_url": updated_data.get("google_lookup_url", current.get("google_lookup_url", "")), "verification_status": override.get("verification_status", current.get("verification_status", "in_process")), "eligibility_status": override.get("eligibility_status", current.get("eligibility_status", "NEEDS_REVIEW")), "data_json": self._json(updated_data), "updated_at": now}
        with self.engine.begin() as conn:
            conn.execute(text("UPDATE admin_venues SET display_name=:display_name, city=:city, state=:state, area=:area, google_place_id=:google_place_id, official_url=:official_url, google_lookup_url=:google_lookup_url, verification_status=:verification_status, eligibility_status=:eligibility_status, data_json=:data_json, updated_at=:updated_at WHERE venue_id=:id"), sql_values)
            conn.execute(text("INSERT INTO admin_audit_log (id,venue_id,action,actor,before_json,after_json,created_at) VALUES (:id,:venue_id,:action,:actor,:before_json,:after_json,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "action": "manual_update", "actor": actor, "before_json": self._json(before), "after_json": self._json(override), "created_at": now})
            conn.execute(text("INSERT INTO venue_change_history (id,venue_id,action,before_json,after_json,actor,created_at) VALUES (:id,:venue_id,:action,:before_json,:after_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "action": "draft_update", "actor": actor, "before_json": self._json(before), "after_json": self._json(updated_data), "created_at": now})
            for key, value in clean.items():
                if key in {"published", "override_reason"}:
                    continue
                conn.execute(text("INSERT INTO venue_manual_overrides (id,venue_id,field_name,automatic_value,manual_value,effective_value,reason,edited_by,edited_at) VALUES (:id,:venue_id,:field_name,:automatic_value,:manual_value,:effective_value,:reason,:edited_by,:edited_at)"), {"id": uuid4().hex, "venue_id": venue_id, "field_name": key, "automatic_value": self._json(current.get(key)), "manual_value": self._json(value), "effective_value": self._json(value), "reason": str(clean.get("override_reason", "")), "edited_by": actor, "edited_at": now})
            sync_normalized_venue(conn, updated_data, now)
        return self.get_venue(venue_id)

    def _transition(self, venue_id: str, status: str, actor: str = "admin") -> dict[str, Any] | None:
        if status not in {"DRAFT", "PUBLISHED", "UNPUBLISHED", "ARCHIVED"}:
            raise ValueError("Invalid publication status")
        current = self.get_venue(venue_id)
        if not current:
            return None
        if status == "PUBLISHED" and current.get("eligibility_status") not in {"AUTO_APPROVED", "PROVISIONAL"}:
            raise ValueError("Venue must be eligible before publishing")
        now = self._now()
        with self.engine.begin() as conn:
            latest = conn.execute(text("SELECT COALESCE(MAX(version_number),0) FROM venue_publication_versions WHERE venue_id=:venue_id"), {"venue_id": venue_id}).scalar() or 0
            if status == "PUBLISHED":
                conn.execute(text("UPDATE admin_venues SET publication_status='PUBLISHED', updated_at=:now WHERE venue_id=:venue_id"), {"now": now, "venue_id": venue_id})
                conn.execute(text("INSERT INTO venue_publication_versions (version_id,venue_id,version_number,status,snapshot_json,created_by,created_at) VALUES (:id,:venue_id,:version_number,'PUBLISHED',:snapshot_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "version_number": int(latest) + 1, "snapshot_json": self._json({**current, "publication_status": "PUBLISHED"}), "actor": actor, "created_at": now})
            else:
                conn.execute(text("UPDATE admin_venues SET publication_status=:status, updated_at=:now WHERE venue_id=:venue_id"), {"status": status, "now": now, "venue_id": venue_id})
            conn.execute(text("INSERT INTO venue_change_history (id,venue_id,action,before_json,after_json,actor,created_at) VALUES (:id,:venue_id,:action,:before_json,:after_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "action": status.lower(), "before_json": self._json(current), "after_json": self._json({**current, "publication_status": status}), "actor": actor, "created_at": now})
        return self.get_venue(venue_id)

    def publish_venue(self, venue_id: str, actor: str = "admin") -> dict[str, Any] | None:
        return self._transition(venue_id, "PUBLISHED", actor)

    def unpublish_venue(self, venue_id: str, actor: str = "admin") -> dict[str, Any] | None:
        return self._transition(venue_id, "UNPUBLISHED", actor)

    def archive_venue(self, venue_id: str, actor: str = "admin") -> dict[str, Any] | None:
        return self._transition(venue_id, "ARCHIVED", actor)

    def restore_venue(self, venue_id: str, actor: str = "admin") -> dict[str, Any] | None:
        return self._transition(venue_id, "DRAFT", actor)

    def history(self, venue_id: str) -> list[dict[str, Any]]:
        with self.engine.connect() as conn:
            return [dict(row) for row in conn.execute(text("SELECT * FROM venue_change_history WHERE venue_id=:venue_id ORDER BY created_at DESC"), {"venue_id": venue_id}).mappings().all()]

    def media(self, venue_id: str) -> list[dict[str, Any]]:
        with self.engine.connect() as conn:
            return [dict(row) for row in conn.execute(text("SELECT * FROM venue_media WHERE venue_id=:venue_id ORDER BY is_cover DESC, sort_order ASC, created_at ASC"), {"venue_id": venue_id}).mappings().all()]

    def media_library(self, search: str = "", page: int = 1, page_size: int = 60) -> dict[str, Any]:
        """Return owned media with venue context for the admin library."""
        page = max(1, page)
        page_size = min(200, max(1, page_size))
        offset = (page - 1) * page_size
        where = "1=1"
        params: dict[str, Any] = {"limit": page_size, "offset": offset}
        if search:
            where = "(m.alt_text LIKE :search OR v.display_name LIKE :search OR v.area LIKE :search)"
            params["search"] = f"%{search}%"
        with self.engine.connect() as conn:
            total = conn.execute(text(f"SELECT COUNT(*) FROM venue_media m JOIN admin_venues v ON v.venue_id=m.venue_id WHERE {where}"), params).scalar() or 0
            rows = conn.execute(text(f"""
                SELECT m.*, v.display_name AS venue_name, v.area, v.city,
                       v.publication_status, v.verification_status
                FROM venue_media m JOIN admin_venues v ON v.venue_id=m.venue_id
                WHERE {where} ORDER BY m.is_cover DESC, m.created_at DESC
                LIMIT :limit OFFSET :offset
            """), params).mappings().all()
        return {"items": [dict(row) for row in rows], "page": page, "page_size": page_size, "total": total}

    def add_media(self, venue_id: str, url: str, media_type: str, source_type: str, rights_status: str, alt_text: str, actor: str) -> dict[str, Any]:
        if not self.get_venue(venue_id):
            raise ValueError("Venue not found")
        now = self._now()
        media_id = uuid4().hex
        with self.engine.begin() as conn:
            next_order = conn.execute(text("SELECT COALESCE(MAX(sort_order), -1) + 1 FROM venue_media WHERE venue_id=:venue_id"), {"venue_id": venue_id}).scalar() or 0
            conn.execute(text("INSERT INTO venue_media (media_id,venue_id,url,media_type,source_type,rights_status,alt_text,sort_order,is_cover,created_by,created_at) VALUES (:media_id,:venue_id,:url,:media_type,:source_type,:rights_status,:alt_text,:sort_order,0,:created_by,:created_at)"), {"media_id": media_id, "venue_id": venue_id, "url": url, "media_type": media_type, "source_type": source_type, "rights_status": rights_status, "alt_text": alt_text, "sort_order": int(next_order), "created_by": actor, "created_at": now})
            conn.execute(text("INSERT INTO venue_change_history (id,venue_id,action,after_json,actor,created_at) VALUES (:id,:venue_id,'media_added',:after_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "after_json": self._json({"url": url, "rights_status": rights_status, "alt_text": alt_text}), "actor": actor, "created_at": now})
        return next(item for item in self.media(venue_id) if item["media_id"] == media_id)

    def remove_media(self, venue_id: str, media_id: str, actor: str) -> bool:
        with self.engine.begin() as conn:
            result = conn.execute(text("DELETE FROM venue_media WHERE venue_id=:venue_id AND media_id=:media_id"), {"venue_id": venue_id, "media_id": media_id})
            if not result.rowcount:
                return False
            conn.execute(text("INSERT INTO venue_change_history (id,venue_id,action,after_json,actor,created_at) VALUES (:id,:venue_id,'media_removed',:after_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "after_json": self._json({"media_id": media_id}), "actor": actor, "created_at": self._now()})
        return True

    def set_cover_media(self, venue_id: str, media_id: str, actor: str) -> bool:
        with self.engine.begin() as conn:
            exists = conn.execute(text("SELECT 1 FROM venue_media WHERE venue_id=:venue_id AND media_id=:media_id"), {"venue_id": venue_id, "media_id": media_id}).fetchone()
            if not exists:
                return False
            conn.execute(text("UPDATE venue_media SET is_cover=0 WHERE venue_id=:venue_id"), {"venue_id": venue_id})
            conn.execute(text("UPDATE venue_media SET is_cover=1 WHERE venue_id=:venue_id AND media_id=:media_id"), {"venue_id": venue_id, "media_id": media_id})
            conn.execute(text("INSERT INTO venue_change_history (id,venue_id,action,after_json,actor,created_at) VALUES (:id,:venue_id,'cover_changed',:after_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "after_json": self._json({"media_id": media_id}), "actor": actor, "created_at": self._now()})
        return True

    def update_media_layout(self, venue_id: str, media_id: str, actor: str, *, sort_order: int | None = None, crop: dict[str, Any] | None = None) -> bool:
        with self.engine.begin() as conn:
            exists = conn.execute(text("SELECT 1 FROM venue_media WHERE venue_id=:venue_id AND media_id=:media_id"), {"venue_id": venue_id, "media_id": media_id}).fetchone()
            if not exists:
                return False
            if sort_order is not None:
                conn.execute(text("UPDATE venue_media SET sort_order=:sort_order WHERE venue_id=:venue_id AND media_id=:media_id"), {"venue_id": venue_id, "media_id": media_id, "sort_order": max(0, sort_order)})
            if crop is not None:
                conn.execute(text("UPDATE venue_media SET crop_json=:crop_json WHERE venue_id=:venue_id AND media_id=:media_id"), {"venue_id": venue_id, "media_id": media_id, "crop_json": self._json(crop)})
            conn.execute(text("INSERT INTO venue_change_history (id,venue_id,action,after_json,actor,created_at) VALUES (:id,:venue_id,'media_layout_changed',:after_json,:actor,:created_at)"), {"id": uuid4().hex, "venue_id": venue_id, "after_json": self._json({"media_id": media_id, "sort_order": sort_order, "crop": crop}), "actor": actor, "created_at": self._now()})
        return True

    def users(self, search: str = "", status: str = "") -> list[dict[str, Any]]:
        with self.engine.connect() as conn:
            clauses = ["1=1"]
            params: dict[str, Any] = {}
            if search:
                clauses.append("(email LIKE :search OR display_name LIKE :search OR user_id LIKE :search)")
                params["search"] = f"%{search}%"
            if status:
                clauses.append("status=:status")
                params["status"] = status
            rows = conn.execute(text(f"SELECT user_id,email,display_name,role,status,last_seen_at,created_at FROM admin_users WHERE {' AND '.join(clauses)} ORDER BY created_at DESC"), params).mappings().all()
            legacy: list[dict[str, Any]] = []
            try:
                legacy_rows = conn.execute(text("SELECT id,email,created_at FROM users ORDER BY created_at DESC")).mappings().all()
                legacy = [{"user_id": row.get("id"), "email": row.get("email"), "display_name": "", "role": "user", "status": "active", "last_seen_at": None, "created_at": row.get("created_at")} for row in legacy_rows]
            except Exception:
                legacy = []
        combined = [dict(row) for row in rows]
        known = {item.get("user_id") for item in combined}
        return combined + [item for item in legacy if item.get("user_id") not in known and (not status or item.get("status") == status) and (not search or search.lower() in str(item).lower())]

    def get_setting(self, key: str, default: str = "") -> str:
        with self.engine.connect() as conn:
            value = conn.execute(text("SELECT setting_value FROM admin_settings WHERE setting_key=:key"), {"key": key}).scalar()
        return str(value) if value is not None else default

    def set_setting(self, key: str, value: str, actor: str) -> None:
        now = self._now()
        with self.engine.begin() as conn:
            existing = conn.execute(text("SELECT 1 FROM admin_settings WHERE setting_key=:key"), {"key": key}).first()
            if existing:
                conn.execute(text("UPDATE admin_settings SET setting_value=:value,updated_by=:actor,updated_at=:now WHERE setting_key=:key"), {"key": key, "value": value, "actor": actor, "now": now})
            else:
                conn.execute(text("INSERT INTO admin_settings(setting_key,setting_value,updated_by,updated_at) VALUES (:key,:value,:actor,:now)"), {"key": key, "value": value, "actor": actor, "now": now})

    def setting_snapshot(self) -> dict[str, str]:
        with self.engine.connect() as conn:
            rows = conn.execute(text("SELECT setting_key,setting_value FROM admin_settings")).all()
        return {str(row[0]): str(row[1]) for row in rows}

    def create_pipeline_run(self, source_dir: str | None = None) -> dict[str, Any]:
        run_id = f"run-{uuid4().hex}"
        now = self._now()
        with self.engine.begin() as conn:
            conn.execute(text("INSERT INTO venue_pipeline_runs (run_id,status,source_dir,started_at) VALUES (:run_id,'QUEUED',:source_dir,:started_at)"), {"run_id": run_id, "source_dir": source_dir or str(DEFAULT_SOURCE), "started_at": now})
        return self.pipeline_run(run_id) or {"run_id": run_id, "status": "QUEUED"}

    def update_pipeline_run(self, run_id: str, status: str, *, summary: dict[str, Any] | None = None, error_count: int = 0) -> dict[str, Any] | None:
        finished_at = self._now() if status in {"SUCCEEDED", "FAILED", "CANCELLED"} else None
        summary = summary or {}
        with self.engine.begin() as conn:
            result = conn.execute(text("UPDATE venue_pipeline_runs SET status=:status, source_rows=:source_rows, inserted_rows=:inserted_rows, updated_rows=:updated_rows, error_count=:error_count, finished_at=:finished_at, summary_json=:summary_json WHERE run_id=:run_id"), {"run_id": run_id, "status": status, "source_rows": int(summary.get("source_rows", summary.get("processed", 0)) or 0), "inserted_rows": int(summary.get("inserted", 0) or 0), "updated_rows": int(summary.get("updated", 0) or 0), "error_count": error_count, "finished_at": finished_at, "summary_json": self._json(summary)})
            if result.rowcount == 0:
                return None
        return self.pipeline_run(run_id)

    def pipeline_run(self, run_id: str) -> dict[str, Any] | None:
        with self.engine.connect() as conn:
            row = conn.execute(text("SELECT * FROM venue_pipeline_runs WHERE run_id=:run_id"), {"run_id": run_id}).mappings().first()
        if not row:
            return None
        item = dict(row)
        raw = item.pop("summary_json", None)
        item["summary"] = json.loads(raw) if raw else {}
        return item

    def pipeline_runs(self, limit: int = 50) -> list[dict[str, Any]]:
        with self.engine.connect() as conn:
            rows = conn.execute(text("SELECT * FROM venue_pipeline_runs ORDER BY started_at DESC LIMIT :limit"), {"limit": max(1, min(limit, 200))}).mappings().all()
        items = []
        for row in rows:
            item = dict(row)
            raw = item.pop("summary_json", None)
            item["summary"] = json.loads(raw) if raw else {}
            items.append(item)
        return items

    def public_venues(self, query: str = "", city: str = "", include_needs_review: bool = False) -> list[dict[str, Any]]:
        """Return app-safe venue cards from the admin source of truth.

        Google display fields are intentionally not copied here. The mobile/web
        client uses google_place_id with its live Places adapter for those fields.
        """
        with self.engine.connect() as conn:
            eligibility = "('AUTO_APPROVED','PROVISIONAL','NEEDS_REVIEW')" if include_needs_review else "('AUTO_APPROVED','PROVISIONAL')"
            publication = "('PUBLISHED','DRAFT')" if include_needs_review else "('PUBLISHED')"
            rows = conn.execute(text(f"SELECT * FROM admin_venues WHERE publication_status IN {publication} AND eligibility_status IN {eligibility} AND verification_status <> 'rejected' ORDER BY area, display_name")).fetchall()
        items: list[dict[str, Any]] = []
        for row in rows:
            item = self._row(row)
            if city and city.lower() not in {str(item.get("city", "")).lower(), "new delhi"}:
                continue
            search_text = " ".join(str(item.get(key, "")) for key in ("display_name", "area", "city", "admission_status"))
            if query and query.lower() not in search_text.lower():
                continue
            override = item.get("admin_override") or {}
            identity = (item.get("venue_identity") or [{}])[0]
            items.append({
                "id": item.get("venue_id"), "name": override.get("display_name") or item.get("display_name") or identity.get("official_name") or "Live Google venue",
                "type": item.get("discovered_type", "venue"), "area": item.get("area") or item.get("city"), "city": item.get("city"),
                "google_place_id": item.get("google_place_id"), "google_lookup_url": item.get("google_lookup_url"),
                "official_url": item.get("official_url"), "tags": override.get("manual_tags", [tag.get("tag") for tag in item.get("venue_tags", [])]),
                "scores": {row.get("factor"): row.get("score") for row in item.get("venue_scores", [])}, "verification_status": item.get("verification_status"),
                "manual_vibe": override.get("manual_vibe", ""), "published": item.get("publication_status") == "PUBLISHED",
                "eligibility_status": item.get("eligibility_status"), "publication_status": item.get("publication_status"),
                "selectable": item.get("eligibility_status") in {"AUTO_APPROVED", "PROVISIONAL"} and item.get("publication_status") == "PUBLISHED",
                "admission_badge": "Needs review" if item.get("eligibility_status") == "NEEDS_REVIEW" else "Admitted",
                "updated_at": item.get("updated_at"), "photos": override.get("photos", []), "dto_version": "venue-dto-v1",
            })
        return items

    def export_xlsx(self) -> BytesIO:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill

        def excel_value(value: Any) -> Any:
            # Google/source payloads can contain binary-like or control
            # characters; Excel rejects those even when they are inside a
            # JSON text column. Keep the export human-readable and lossless
            # enough for the source URL/evidence review.
            if isinstance(value, (bytes, bytearray, memoryview)):
                value = repr(bytes(value))
            if isinstance(value, str):
                value = "".join(char for char in value if char in "\t\n\r" or ord(char) >= 32)
            return value

        # The catalogue contains tens of thousands of normalized rows.  A
        # normal workbook keeps every cell object in memory and can appear to
        # hang on Windows; write-only mode streams rows safely instead.
        workbook = Workbook(write_only=True)
        sheet_specs = [
            ("Venues", "SELECT * FROM admin_venues ORDER BY area, display_name"),
            ("Venue_Identity", "SELECT * FROM venue_identity ORDER BY official_name"),
            ("Venue_Scores", "SELECT * FROM venue_scores ORDER BY venue_id, factor"),
            ("Venue_Tags", "SELECT * FROM venue_tags ORDER BY venue_id, tag"),
            ("Venue_Sources", "SELECT * FROM venue_sources ORDER BY venue_id, field_name"),
            ("Venue_Time_Slots", "SELECT * FROM venue_time_slots ORDER BY venue_id, day_of_week, start_time"),
            ("Venue_Metros", "SELECT * FROM venue_metros ORDER BY venue_id, distance_km"),
            ("Users", "SELECT * FROM admin_users ORDER BY created_at DESC"),
            ("Audit_Log", "SELECT * FROM admin_audit_log ORDER BY created_at DESC"),
        ]
        with self.engine.connect() as conn:
            for name, query in sheet_specs:
                sheet = workbook.create_sheet(name)
                rows = conn.execute(text(query)).mappings().all()
                columns = list(rows[0].keys()) if rows else ["status"]
                header_cells = []
                for column in columns:
                    cell = WriteOnlyCell(sheet, value=column)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="17324D")
                    header_cells.append(cell)
                sheet.append(header_cells)
                for row in rows:
                    sheet.append([excel_value(row.get(column)) for column in columns])
                # Freeze panes/auto-filter are intentionally omitted in
                # streaming mode; consumers can still filter the tab and the
                # export remains valid even for very large sheets.
        buffer = BytesIO(); workbook.save(buffer); buffer.seek(0); return buffer


store = AdminStore()

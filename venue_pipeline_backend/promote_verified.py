from __future__ import annotations

import csv
import argparse
from pathlib import Path

ROOT = Path(__file__).resolve().parent
QUEUE = ROOT / "data_market_ready" / "SHUFFL_Delhi_NCR_Client_Verification_Final.xlsx"
OUT = ROOT / "config" / "official_verification.csv"


def load_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["Auto_Review_Queue"]
        values = list(sheet.iter_rows(min_row=3, values_only=True))
        headers = [str(value or "") for value in values[0]]
        return [dict(zip(headers, ["" if value is None else str(value) for value in row])) for row in values[1:] if any(value is not None for value in row)]
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("file", type=Path, nargs="?", default=QUEUE)
    args = parser.parse_args()
    rows = load_rows(args.file)
    fields = ["google_place_id", "verification_status", "name", "address", "area", "city", "latitude", "longitude", "venue_type", "official_website", "official_instagram", "official_facebook", "menu_url", "booking_url", "official_website_checked", "official_social_checked", "hours_confirmed", "scene_confirmed", "hours_json", "reviewer_name", "reviewed_at", "notes"]
    promoted = []
    for row in rows:
        status = (row.get("human_status") or "").strip().lower()
        if status != "verified":
            continue
        socials = [value.strip() for value in (row.get("social_links") or "").strip("[]").replace("'", "").split(",") if value.strip()]
        promoted.append({"google_place_id": row.get("google_place_id", ""), "verification_status": "verified", "name": row.get("official_name", ""), "address": row.get("official_address", ""), "area": row.get("discovered_area", ""), "city": row.get("discovered_city", ""), "latitude": "", "longitude": "", "venue_type": row.get("discovered_type", ""), "official_website": row.get("official_website", ""), "official_instagram": next((value for value in socials if "instagram.com" in value), ""), "official_facebook": next((value for value in socials if "facebook.com" in value), ""), "menu_url": "", "booking_url": "", "official_website_checked": "true", "official_social_checked": "true" if socials else "false", "hours_confirmed": "true" if row.get("hours_json") else "false", "scene_confirmed": "true", "hours_json": row.get("hours_json", ""), "reviewer_name": row.get("reviewer_name", ""), "reviewed_at": row.get("reviewed_at", ""), "notes": row.get("human_notes", "")})
    with OUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(promoted)
    print(f"Promoted verified venues: {len(promoted)} -> {OUT}")


if __name__ == "__main__":
    main()

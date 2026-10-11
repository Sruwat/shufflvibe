from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

ALLOWED = {"verified", "rejected", "in_process", "update_required"}
ROOT = Path(__file__).resolve().parent


def load_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["human_pass"]
        values = list(sheet.iter_rows(min_row=3, values_only=True))
        headers = [str(value or "") for value in values[0]]
        return [dict(zip(headers, ["" if value is None else str(value) for value in row])) for row in values[1:] if any(value is not None for value in row)]
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("file", type=Path)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data_market_ready_v2")
    args = parser.parse_args()
    rows = load_rows(args.file)
    invalid = [row.get("venue_id", "") for row in rows if (row.get("human_status") or "").strip() not in ALLOWED]
    if invalid:
        raise SystemExit(f"Invalid or blank human_status in {len(invalid)} rows. Allowed: {sorted(ALLOWED)}")
    verified = [{"venue_id": row.get("venue_id", ""), "google_place_id": row.get("google_place_id", ""), "google_lookup_url": row.get("google_lookup_url", ""), "fit_out_age": row.get("fit_out_age", ""), "operator_class": row.get("operator_class", ""), "instagram_schedule": row.get("instagram_schedule", ""), "human_status": "verified", "human_notes": row.get("human_notes", ""), "status": "active_live_google"} for row in rows if row.get("human_status", "").strip() == "verified"]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "verified_venue_index.json").write_text(json.dumps({"venues": verified}, ensure_ascii=False, indent=2), encoding="utf-8")
    with (args.output_dir / "verified_venue_index.csv").open("w", encoding="utf-8", newline="") as handle:
        fields = list(verified[0].keys()) if verified else ["venue_id", "google_place_id", "google_lookup_url", "fit_out_age", "operator_class", "instagram_schedule", "human_status", "human_notes", "status"]
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(verified)
    print(f"Verified active venues: {len(verified)}")
    print(f"Review/rejected venues: {len(rows) - len(verified)}")


if __name__ == "__main__":
    main()
